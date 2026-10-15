import logging
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

from app.core.config import settings
from app.core.errors import ErrorCode, raise_app_error
from app.core.logging_utils import log_payload, new_trace_id
from app.db.store import Store
from app.llm.factory import create_llm_client


logger = logging.getLogger("uvicorn.error")


class ResumeService:
    allowed_extensions = {".pdf", ".txt", ".md", ".docx"}

    def __init__(self) -> None:
        self.store = Store()

    async def parse_upload(
        self, *, user_id: int, file: UploadFile, model_provider: str = "", model_name: str = ""
    ) -> dict[str, Any]:
        trace_id = new_trace_id("resume")
        # 3.3：先查扩展名，避免恶意大文件占内存
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in self.allowed_extensions:
            logger.warning(
                "[trace_id=%s] Unsupported resume file extension: filename=%s suffix=%s user_id=%s",
                trace_id,
                file.filename or "resume",
                suffix or "-",
                user_id,
            )
            raise_app_error(400, ErrorCode.RESUME_FILE_TYPE_UNSUPPORTED, "仅支持 pdf/txt/md/docx 简历文件", trace_id=trace_id)

        # 3.3：先读 max+1 字节，超限立即拒绝
        content = await file.read(settings.max_upload_bytes + 1)
        if len(content) > settings.max_upload_bytes:
            raise_app_error(413, ErrorCode.RESUME_TOO_LARGE, f"简历文件过大（上限 5MB）", trace_id=trace_id)

        logger.info(
            "[trace_id=%s] Resume upload received: filename=%s suffix=%s size_bytes=%s user_id=%s",
            trace_id,
            file.filename or "resume",
            suffix or "-",
            len(content),
            user_id,
        )
        resume_text = self._extract_text(file.filename or "resume", suffix, content, trace_id=trace_id)
        log_payload(
            logger,
            logging.INFO,
            f"Resume extracted text: filename={file.filename or 'resume'} suffix={suffix or '-'}",
            resume_text,
            trace_id=trace_id,
        )
        if len(resume_text.strip()) < 20:
            logger.warning(
                "[trace_id=%s] Resume text too short after extraction: filename=%s text_length=%s user_id=%s",
                trace_id,
                file.filename or "resume",
                len(resume_text.strip()),
                user_id,
            )
            raise_app_error(400, ErrorCode.RESUME_TEXT_TOO_SHORT, "简历内容过短，无法解析", trace_id=trace_id)

        resolved_provider = (model_provider or settings.llm_provider or "").strip()
        resolved_model_name = (model_name or settings.llm_model or "").strip()
        logger.info(
            "[trace_id=%s] Resume parse requested: provider=%s model=%s filename=%s user_id=%s",
            trace_id,
            resolved_provider or "-",
            resolved_model_name or "-",
            file.filename or "resume",
            user_id,
        )
        parsed = self._parse_resume_text(
            resume_text,
            model_provider=resolved_provider,
            model_name=resolved_model_name,
            trace_id=trace_id,
        )
        log_payload(
            logger,
            logging.INFO,
            f"Resume structured summary: filename={file.filename or 'resume'} provider={resolved_provider or '-'}",
            parsed,
            trace_id=trace_id,
        )
        parsed["filename"] = file.filename or "resume"
        parsed["resume_text"] = resume_text
        parsed["model_provider"] = resolved_provider
        parsed["model_name"] = resolved_model_name
        saved = self.store.create_resume(user_id=user_id, filename=parsed["filename"], payload=parsed)
        log_payload(
            logger,
            logging.INFO,
            f"Resume saved payload: resume_id={saved.get('id', '')} filename={parsed['filename']}",
            saved,
            trace_id=trace_id,
        )
        return saved

    def list_resumes(self, user_id: int) -> list[dict[str, Any]]:
        return self.store.list_resumes(user_id)

    def get_resume(self, user_id: int, resume_id: str) -> dict[str, Any]:
        resume = self.store.get_resume(user_id, resume_id)
        if not resume:
            raise_app_error(404, ErrorCode.RESUME_NOT_FOUND, "简历不存在")
        return resume

    def delete_resume(self, user_id: int, resume_id: str) -> None:
        if not self.store.delete_resume(user_id, resume_id):
            raise_app_error(404, ErrorCode.RESUME_NOT_FOUND, "简历不存在")

    def _extract_text(self, filename: str, suffix: str, content: bytes, *, trace_id: str) -> str:
        if suffix in {".txt", ".md"}:
            logger.info("[trace_id=%s] Decoding text resume directly: filename=%s", trace_id, filename)
            return content.decode("utf-8", errors="ignore")
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            logger.info("[trace_id=%s] Extracting PDF text: filename=%s page_count=%s", trace_id, filename, len(reader.pages))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix == ".docx":
            document = Document(BytesIO(content))
            logger.info(
                "[trace_id=%s] Extracting DOCX text: filename=%s paragraph_count=%s",
                trace_id,
                filename,
                len(document.paragraphs),
            )
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        raise_app_error(400, ErrorCode.RESUME_FILE_TYPE_UNSUPPORTED, f"不支持的文件类型: {filename}", trace_id=trace_id)

    def _parse_resume_text(
        self, resume_text: str, *, model_provider: str, model_name: str, trace_id: str
    ) -> dict[str, Any]:
        try:
            return self._parse_with_llm(
                resume_text,
                model_provider=model_provider,
                model_name=model_name,
                trace_id=trace_id,
            )
        except ValueError as exc:
            message = str(exc)
            if message in {"empty response", "ANTHROPIC_API_KEY 未配置"} or "ANTHROPIC" in message:
                logger.exception(
                    "[trace_id=%s] Resume parse configuration error: provider=%s model=%s",
                    trace_id,
                    model_provider or "-",
                    model_name or "-",
                )
                raise_app_error(500, ErrorCode.MODEL_CONFIG_INVALID, f"模型配置错误: {exc}", trace_id=trace_id)
            logger.exception(
                "[trace_id=%s] Resume parse failed while parsing model JSON: provider=%s model=%s",
                trace_id,
                model_provider or "-",
                model_name or "-",
            )
            raise_app_error(502, ErrorCode.RESUME_PARSE_FAILED, f"简历解析失败: {exc}", trace_id=trace_id)
        except HTTPException:
            raise
        except Exception as exc:
            logger.exception(
                "[trace_id=%s] Resume parse failed unexpectedly: provider=%s model=%s",
                trace_id,
                model_provider or "-",
                model_name or "-",
            )
            raise_app_error(502, ErrorCode.RESUME_PARSE_FAILED, f"简历解析失败: {exc}", trace_id=trace_id)

    def _parse_with_llm(self, resume_text: str, *, model_provider: str, model_name: str, trace_id: str) -> dict[str, Any]:
        client = create_llm_client(model_provider, model_name)
        prompt = f"""
你是一个教学版简历解析助手，请从下面简历中抽取结构化信息，并返回 JSON。

简历内容:
{resume_text}

返回格式:
{{
  "personal_info": "一句话个人信息摘要",
  "highlights": ["亮点标签"],
  "technical_skills": ["技能1", "技能2"],
  "work_experience": [{{"name": "公司/经历名", "detail": "内容"}}],
  "project_experience": [{{"name": "项目名", "detail": "内容"}}]
}}

要求：
1. 只返回 JSON。
2. `highlights` 和 `technical_skills` 最多 8 个。
3. 缺失字段返回空数组或空字符串。
4. 不要过度润色，保持教学版简洁风格。
"""
        try:
            payload = client.generate_json(
                prompt,
                "你是结构化输出助手，只能返回 JSON。",
                operation_name="resume_parse",
                trace_id=trace_id,
            )
            normalized = {
                "personal_info": str(payload.get("personal_info", "")).strip(),
                "highlights": self._to_clean_list(payload.get("highlights", []), limit=8),
                "technical_skills": self._to_clean_list(payload.get("technical_skills", []), limit=8),
                "work_experience": self._normalize_entries(payload.get("work_experience", [])),
                "project_experience": self._normalize_entries(payload.get("project_experience", [])),
                "parse_source": "llm",
                "parse_degraded": False,
            }
            log_payload(logger, logging.INFO, "Resume parsed with LLM and normalized", normalized, trace_id=trace_id)
            return normalized
        except Exception:
            logger.exception(
                "[trace_id=%s] Resume LLM parse failed, using rule fallback: provider=%s model=%s",
                trace_id,
                model_provider or "-",
                model_name or "-",
            )
            parsed = self._parse_with_rules(resume_text)
            parsed["parse_source"] = "rule_fallback"
            parsed["parse_degraded"] = True
            log_payload(logger, logging.INFO, "Resume parsed with rule fallback", parsed, trace_id=trace_id)
            return parsed

    def _parse_with_rules(self, resume_text: str) -> dict[str, Any]:
        lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
        personal_info = "；".join(lines[:3])[:240]

        sections = self._split_sections(resume_text)
        work_experience = self._make_entries(sections.get("work", []), prefix="工作经历")
        project_experience = self._make_entries(sections.get("project", []), prefix="项目经历")
        technical_skills = self._extract_skills(resume_text)
        highlights = technical_skills[:4] or ["项目表达", "技术取舍", "问题排查"]

        return {
            "personal_info": personal_info,
            "highlights": highlights,
            "technical_skills": technical_skills,
            "work_experience": work_experience,
            "project_experience": project_experience,
        }

    @staticmethod
    def _to_clean_list(items: Any, *, limit: int) -> list[str]:
        if not isinstance(items, list):
            return []
        result: list[str] = []
        for item in items:
            text = str(item).strip()
            if text and text not in result:
                result.append(text)
            if len(result) >= limit:
                break
        return result

    @staticmethod
    def _normalize_entries(items: Any) -> list[dict[str, str]]:
        if not isinstance(items, list):
            return []
        result: list[dict[str, str]] = []
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                continue
            detail = str(item.get("detail", "")).strip()
            if not detail:
                continue
            result.append({"name": str(item.get("name", f"经历 {index}")).strip() or f"经历 {index}", "detail": detail})
        return result[:6]

    @staticmethod
    def _split_sections(resume_text: str) -> dict[str, list[str]]:
        work_markers = ("工作经历", "实习经历", "experience")
        project_markers = ("项目经历", "项目经验", "project")
        work_lines: list[str] = []
        project_lines: list[str] = []
        current = None
        for raw_line in resume_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            lower = line.lower()
            if any(marker in line or marker in lower for marker in work_markers):
                current = "work"
                continue
            if any(marker in line or marker in lower for marker in project_markers):
                current = "project"
                continue
            if current == "work":
                work_lines.append(line)
            elif current == "project":
                project_lines.append(line)
        return {"work": work_lines, "project": project_lines}

    @staticmethod
    def _make_entries(lines: list[str], *, prefix: str) -> list[dict[str, str]]:
        if not lines:
            return []
        chunks: list[list[str]] = []
        current: list[str] = []
        for line in lines:
            if re.match(r"^[-•\d一二三四五六七八九十]", line) and current:
                chunks.append(current)
                current = [line]
            else:
                current.append(line)
        if current:
            chunks.append(current)
        entries: list[dict[str, str]] = []
        for index, chunk in enumerate(chunks[:4], start=1):
            first = chunk[0][:40]
            entries.append({"name": f"{prefix} {index}: {first}", "detail": "\n".join(chunk)[:1200]})
        return entries

    @staticmethod
    def _extract_skills(resume_text: str) -> list[str]:
        candidates = [
            "Python",
            "Java",
            "Go",
            "SQL",
            "MySQL",
            "PostgreSQL",
            "Redis",
            "Kafka",
            "Docker",
            "Kubernetes",
            "Spring",
            "FastAPI",
            "Flask",
            "React",
            "TypeScript",
            "LLM",
            "Prompt",
            "微服务",
            "性能优化",
            "接口设计",
        ]
        lower = resume_text.lower()
        return [skill for skill in candidates if skill.lower() in lower][:8]
